import streamlit as st
import io
import os
import requests
import xml.etree.ElementTree as ET
from urllib.parse import quote

# AI 클라이언트 (선택적으로 사용)
try:
    from google import genai
except ImportError:
    genai = None

try:
    import anthropic
except ImportError:
    anthropic = None

# PDF / DOCX 파일 읽기
try:
    import fitz  # PyMuPDF
except ImportError:
    fitz = None

try:
    from docx import Document
except ImportError:
    Document = None


# --------------------------------------------------
# 기본 설정
# --------------------------------------------------

st.set_page_config(
    page_title="AI 취업·자소서 분석기",
    page_icon="📄",
    layout="wide"
)

st.title("📄 AI 취업·자소서 분석기")
st.caption("자소서 + 채용공고 + 지원직무를 비교해서 개선 방향을 분석하거나, 새 자소서를 항목별로 생성합니다.")


# --------------------------------------------------
# AI 제공자 / API 키 / 모델
# --------------------------------------------------

with st.sidebar:
    st.header("⚙️ 설정")

    provider = st.radio(
        "🤖 사용할 AI",
        ["Google Gemini", "Anthropic Claude"],
        horizontal=True
    )

    if provider == "Google Gemini":
        api_key = st.text_input(
            "Google Gemini API Key",
            type="password",
            help="API Key는 이 채팅에 보내지 말고 여기 직접 입력해 주세요"
        )
        model_name = st.text_input(
            "Gemini 모델명",
            value="gemini-3.6-flash",
            help="Google AI Studio에서 사용 가능한 정확한 모델명을 입력하세요."
        )
    else:
        api_key = st.text_input(
            "Anthropic Claude API Key",
            type="password",
            help="API Key는 이 채팅에 보내지 말고 여기 직접 입력해 주세요"
        )
        model_name = st.text_input(
            "Claude 모델명",
            value="claude-sonnet-4-6",
            help="Anthropic 콘솔에서 사용 가능한 정확한 모델명을 입력하세요."
        )

    if api_key:
        st.success("✅ API 키가 입력되었습니다!")
    else:
        st.warning("⚠️ API 키를 입력해 주세요.")

    st.divider()

    st.markdown("""
    ### 현재 버전

    **2차 버전**

    - Gemini / Claude 중 선택 가능
    - 기존 자소서 분석 / AI 수정본 생성
    - **새 자소서 항목별 생성 (신규)**
    - **자소서 항목 선택란 (신규)**
    - 직무 적합도 분석, 강점/약점 분석
    - 원본/수정본 비교, 결과 다운로드 (.txt / .docx)
    - 글자 수 실시간 카운터

    이후에는

    - 실제 채용공고 검색
    - 회사 뉴스 분석 / 채용시장 분석
    - 경쟁 지원자 포지셔닝

    을 추가할 수 있습니다.
    """)

st.subheader("🏢 기업 심화 정보 & 벤치마킹 (선택사항)")

company_philosophy = st.text_area(
    "💡 기업의 철학, 고유 특징 및 내부 암묵지/우선 고려사항",
    placeholder="예: 고객 중심 사고 강조, 실패를 두려워하지 않는 실험 정신 등",
    height=120
)

reference_resume = st.text_area(
    "🏆 합격자 / 우수 자기소개서 레퍼런스",
    placeholder="비교 분석받고 싶은 합격자의 자기소개서를 붙여넣으세요.",
    height=200
)


# --------------------------------------------------
# 모드 선택 (신규)
# --------------------------------------------------

st.divider()
mode = st.radio(
    "🎯 무엇을 하고 싶으신가요?",
    ["🔍 기존 자소서 분석 및 수정", "✨ 새 자소서 항목별로 생성"],
    horizontal=True
)
st.divider()


# --------------------------------------------------
# 공통 컨텍스트 블록 생성 함수
# --------------------------------------------------

def build_context_blocks(philosophy, reference, news_items=None):
    philosophy_block = (
        f"\n[기업 철학 / 암묵지 / 우선 고려사항]\n{philosophy}\n"
        if philosophy and philosophy.strip() else ""
    )

    reference_block = (
        f"\n[합격자/우수 자기소개서 레퍼런스 — 톤·구성 참고용, 그대로 베끼면 안 됨]\n{reference}\n"
        if reference and reference.strip() else ""
    )

    news_block = ""
    if news_items:
        news_lines = "\n".join(
            f"- {n['title']} ({n['source']})" for n in news_items
        )
        news_block = (
            f"\n[해당 직무/업계 최신 뉴스 헤드라인 — 배경 참고용, 사실로 단정하지 말고 참고만 하세요]\n{news_lines}\n"
        )

    return philosophy_block, reference_block, news_block


# --------------------------------------------------
# 공통 AI 호출 함수 (Gemini / Claude 분기)
# --------------------------------------------------

def call_ai(provider, api_key, model, prompt, max_tokens=4096):
    """provider에 따라 Gemini 또는 Claude API를 호출하고 텍스트를 반환합니다."""

    if provider == "Google Gemini":
        if genai is None:
            raise RuntimeError("google-genai 패키지가 설치되어 있지 않습니다. (pip install google-genai)")

        client = genai.Client(api_key=api_key)
        response = client.models.generate_content(
            model=model,
            contents=prompt
        )
        return response.text

    else:  # Anthropic Claude
        if anthropic is None:
            raise RuntimeError("anthropic 패키지가 설치되어 있지 않습니다. (pip install anthropic)")

        client = anthropic.Anthropic(api_key=api_key)
        response = client.messages.create(
            model=model,
            max_tokens=max_tokens,
            messages=[
                {"role": "user", "content": prompt}
            ]
        )
        # Claude 응답은 content가 블록 리스트 형태
        return "".join(
            block.text for block in response.content if hasattr(block, "text")
        )


# --------------------------------------------------
# 파일에서 텍스트 추출
# --------------------------------------------------

def extract_text(uploaded_file):

    file_name = uploaded_file.name.lower()
    file_bytes = uploaded_file.getvalue()

    if file_name.endswith(".txt"):
        return file_bytes.decode("utf-8", errors="ignore")

    elif file_name.endswith(".pdf"):

        if fitz is None:
            st.error("PDF 처리를 위해 PyMuPDF 설치가 필요합니다.")
            return ""

        text = ""
        pdf = fitz.open(stream=file_bytes, filetype="pdf")
        for page in pdf:
            text += page.get_text()
        pdf.close()
        return text

    elif file_name.endswith(".docx"):

        if Document is None:
            st.error("DOCX 처리를 위해 python-docx 설치가 필요합니다.")
            return ""

        document = Document(io.BytesIO(file_bytes))
        text = [paragraph.text for paragraph in document.paragraphs]
        return "\n".join(text)

    else:
        return ""


def count_chars(text):
    """공백 포함 글자 수를 반환합니다."""
    return len(text) if text else 0


def fetch_related_news(keyword, max_items=5):
    """구글 뉴스 RSS(무료, API 키 불필요)에서 키워드 관련 최신 뉴스를 가져옵니다."""

    if not keyword or not keyword.strip():
        return []

    url = f"https://news.google.com/rss/search?q={quote(keyword)}&hl=ko&gl=KR&ceid=KR:ko"

    response = requests.get(url, timeout=10)
    response.raise_for_status()

    root = ET.fromstring(response.content)

    items = []
    for item in root.findall(".//item")[:max_items]:
        title = item.findtext("title") or ""
        link = item.findtext("link") or ""
        pub_date = item.findtext("pubDate") or ""
        source_el = item.find("source")
        source_name = source_el.text if source_el is not None else ""

        items.append({
            "title": title,
            "link": link,
            "pub_date": pub_date,
            "source": source_name
        })

    return items


def build_docx_bytes(title, body_text):
    """간단한 텍스트를 docx 바이트로 변환합니다. python-docx가 없으면 None 반환."""
    if Document is None:
        return None

    doc = Document()
    doc.add_heading(title, level=1)
    for line in body_text.split("\n"):
        doc.add_paragraph(line)

    buf = io.BytesIO()
    doc.save(buf)
    buf.seek(0)
    return buf.getvalue()


# --------------------------------------------------
# AI 분석 함수 (기존)
# --------------------------------------------------

def analyze_resume(
    provider,
    api_key,
    model,
    resume_text,
    company,
    job,
    job_posting,
    profile,
    philosophy,
    reference,
    length_hint,
    news_items=None
):
    length_text = (
        f"- **글자 수 제한**: 공백 포함 **{length_hint}** 분량으로 작성해 주세요."
        if length_hint else ""
    )

    philosophy_block, reference_block, news_block = build_context_blocks(philosophy, reference, news_items)

    prompt = f"""
{length_text}
당신은 한국 취업시장과 채용서류 분석을 전문으로 하는 커리어 컨설턴트입니다.

아래 지원자의 자기소개서와 지원 정보를 분석하세요.

중요한 원칙:

1. 지원자가 실제로 제공하지 않은 경험, 자격증, 성과, 숫자 등을 만들어내지 마세요.
2. 확인되지 않은 채용 통계나 평균 연령을 사실처럼 말하지 마세요.
3. 현재 제공된 자료에서 확인할 수 있는 사실과 AI의 추론/추천을 구분하세요.
4. 자기소개서의 문장 자체뿐 아니라 실제 채용 관점에서 평가하세요.
5. 단순히 "좋습니다/부족합니다"가 아니라 왜 그런지 설명하세요.
6. 채용공고의 요구사항과 자기소개서 내용을 하나씩 대응시키세요.
7. 기업 철학/암묵지가 제공된 경우, 자기소개서가 그 철학과 얼마나 부합하는지도 평가하세요.
8. 합격자 레퍼런스가 제공된 경우, 내용을 베끼라는 뜻이 아니라 구성/톤/강조점 비교 참고용으로만 사용하세요.
9. 최신 뉴스 헤드라인이 제공된 경우, 업계/직무의 최근 분위기를 이해하는 배경 정보로만 참고하고, 자기소개서 평가의 확정적 근거로 사용하지 마세요.

[지원 회사]
{company}

[지원 직무]
{job}

[채용공고]
{job_posting}
{philosophy_block}{reference_block}{news_block}
[지원자 추가 정보 / 스펙]
{profile}

[자기소개서]
{resume_text}

다음 순서로 분석해주세요.

# 1. 자기소개서 핵심 요약
지원자가 어떤 사람으로 보이는지 5줄 이내로 요약

# 2. 지원 직무 적합도
100점 만점의 AI 평가점수를 제시하고,
점수는 실제 합격확률이 아니라 현재 자기소개서와 직무 요구사항의 적합도를 의미한다고 명시

# 3. 채용공고 요구사항 매칭
다음 형식으로 정리

- 요구사항:
- 자기소개서에서 확인되는 근거:
- 충족 정도:
- 부족한 부분:

# 4. 기업 철학 부합도 (정보가 제공된 경우에만)
기업 철학/암묵지와 자기소개서 내용의 부합 정도

# 5. 자기소개서의 강점
실제 채용담당자 입장에서 강점이 될 만한 부분

# 6. 자기소개서의 약점
서류 평가에서 감점될 가능성이 있는 부분

# 7. 빠져 있는 핵심 요소
현재 자기소개서에 없지만 지원 직무에 중요할 가능성이 높은 내용

# 8. 추가하면 좋은 내용
지원자가 실제로 가지고 있을 가능성이 있는 경험을 추측하지 말고,
"이런 경험이 있다면 추가하세요"라는 방식으로 제안

# 9. 문장 개선 포인트
추상적인 표현, 근거 부족, 중복, 지나친 자기칭찬 등을 찾아 설명

# 10. 최우선 개선사항 TOP 5
가장 중요한 것부터 순서대로 제시

# 11. 종합 평가
현재 자소서가 어떤 포지션으로 보이는지 설명
"""

    return call_ai(provider, api_key, model, prompt)


# --------------------------------------------------
# AI 수정본 (기존)
# --------------------------------------------------

def revise_resume(
    provider,
    api_key,
    model,
    resume_text,
    analysis,
    company,
    job,
    length_hint
):
    length_text = (
        f"- **글자 수 제한**: 공백 포함 **{length_hint}** 분량을 지켜서 작성해 주세요."
        if length_hint else ""
    )

    prompt = f"""
당신은 한국 기업 채용 자기소개서 전문 컨설턴트입니다.

아래 자기소개서를 지원 회사와 직무에 맞게 개선하세요.

{length_text}

[회사]
{company}

[직무]
{job}

[기존 자기소개서]
{resume_text}

[기존 분석]
{analysis}

수정 원칙:

1. 기존 지원자의 실제 경험과 사실을 최대한 보존하세요.
2. 없는 경험이나 성과를 새로 만들어내지 마세요.
3. 숫자, 자격증, 수상, 인턴, 프로젝트 등을 임의로 추가하지 마세요.
4. 표현만 더 전문적으로 바꿀 수 있습니다.
5. 지원 직무와 연결되는 부분은 더 명확하게 표현하세요.
6. STAR 방식이 적용될 수 있는 부분은 상황-행동-결과가 드러나도록 개선하세요.
7. 지나치게 AI가 작성한 것처럼 보이는 문장을 피하세요.
8. 원래 지원자의 말투와 경험을 최대한 유지하세요.
9. 글자 수 제한이 주어진 경우 반드시 그 범위 내로 작성하세요.

다음 형식으로 답변하세요.

# AI 수정 자기소개서

전체 수정본을 작성하세요.

# 주요 수정사항

각 수정사항마다:

1. 기존 표현
2. 수정 표현
3. 수정 이유

형식으로 설명하세요.

# 추가하면 좋은 실제 경험

현재 정보만으로 확인할 수 없는 내용은
"해당 경험이 있다면 추가"라고 표시하세요.
"""

    return call_ai(provider, api_key, model, prompt)


# --------------------------------------------------
# AI 자소서 신규 생성 함수 (신규)
# --------------------------------------------------

def generate_resume(
    provider,
    api_key,
    model,
    company,
    job,
    job_posting,
    profile,
    philosophy,
    reference,
    items,
    length_hint,
    news_items=None
):
    length_text = (
        f"- **항목당 글자 수 목표**: 공백 포함 **{length_hint}** 분량으로 작성해 주세요."
        if length_hint else "- 항목당 글자 수 제한은 없지만, 너무 짧거나 장황하지 않게 적절한 분량으로 작성하세요."
    )

    philosophy_block, reference_block, news_block = build_context_blocks(philosophy, reference, news_items)

    items_text = "\n".join(f"{i+1}. {item}" for i, item in enumerate(items))

    prompt = f"""
당신은 한국 채용시장, 특히 병원/기업 자기소개서 작성을 전문으로 하는 15년 경력의 커리어 컨설턴트입니다.
아래 지원자 정보를 바탕으로, 요청된 항목별로 "실제로 서류 합격을 이끌어낼 수 있는 수준"의
완성도 높은 자기소개서를 새로 작성하세요.

절대 원칙 (반드시 지킬 것):

1. 지원자가 [지원자 정보/스펙]에 제공한 사실 범위 안에서만 작성하세요. 없는 경력, 자격증, 수치, 수상 실적을 절대 지어내지 마세요.
2. 제공된 정보가 부족한 항목은, 있는 정보를 최대한 활용해서 작성하되 문장 끝에 "(※ 이 부분은 지원자의 실제 구체적 경험/수치로 반드시 보완 필요)"라고 표시하세요.
3. "최선을 다하겠습니다", "열심히 하겠습니다", "책임감을 가지고" 같은 상투적이고 뻔한 문구의 반복을 피하고, 구체적 상황과 행동이 드러나는 문장으로 쓰세요.
4. 지원자가 준 정보들 중 사용 가능한 사례가 있다면 STAR 기법(상황-행동-결과) 구조로 풀어서 설득력을 높이세요.
5. 기업 철학/암묵지 정보가 제공된 경우, 문장에 자연스럽게(작위적이지 않게) 녹여내세요.
6. 합격자 레퍼런스가 제공된 경우, 표현을 그대로 베끼지 말고 구성 방식과 강조점만 참고하세요.
7. 최신 뉴스 헤드라인이 제공된 경우, 업계/직무 이해를 보여주는 배경으로만 자연스럽게 참고하세요 (사실 단정 금지).
8. 항목들 사이에 내용이 겹치지 않게, 각 항목마다 다른 강점/경험을 부각하세요.
9. AI가 쓴 것 같은 인공적인 느낌을 최소화하고, 실제 사람이 쓴 것 같은 자연스러운 문장으로 작성하세요.

{length_text}

[지원 회사]
{company}

[지원 직무]
{job}

[채용공고]
{job_posting if job_posting else "(제공되지 않음 - 일반적인 해당 직무 기준으로 작성)"}
{philosophy_block}{reference_block}{news_block}
[지원자 정보 / 스펙 / 경험]
{profile}

작성할 자기소개서 항목 (아래 순서대로, 각 항목마다 완성된 답변을 작성하세요):
{items_text}

출력 형식 (반드시 이 형식을 지킬 것):

각 항목마다 아래 형식을 반복하세요.

## [항목 번호]. [항목명]

(여기에 완성된 자기소개서 답변 작성. 문단 구분해서 읽기 쉽게)

**✏️ 작성 포인트**: (이 항목에서 어떤 전략으로 썼는지, 왜 이 구성을 택했는지 1~2줄로 설명)

**💡 더 강력하게 만들려면**: (지원자가 추가로 채워 넣으면 훨씬 좋아질 구체적 정보나 경험 1가지 제안)

모든 항목 작성이 끝나면 마지막에 아래 섹션을 추가하세요.

## 📋 전체 총평
이 지원자의 자소서 세트가 전체적으로 어떤 인상을 주는지, 그리고 서류 통과 가능성을 높이기 위해
가장 먼저 보완해야 할 것 1~2가지를 제시하세요.
"""

    return call_ai(provider, api_key, model, prompt, max_tokens=8192)


# --------------------------------------------------
# 입력 화면 - 공통 (지원 정보 / 채용공고)
# --------------------------------------------------

st.header("1️⃣ 지원 정보 입력")

col1, col2 = st.columns(2)

with col1:
    company = st.text_input("🏢 지원 회사", placeholder="예: 삼성서울병원")

with col2:
    job = st.text_input("💼 지원 직무", placeholder="예: 신규간호사")


st.header("2️⃣ 채용공고")

job_posting = st.text_area(
    "📋 채용공고 내용을 붙여넣으세요"
    + ("" if mode == "🔍 기존 자소서 분석 및 수정" else " (선택사항 — 없어도 생성 가능하지만, 있으면 훨씬 정확해져요)"),
    height=220,
    placeholder="""채용공고의 주요 내용을 그대로 붙여넣으세요.

예:
- 지원자격
- 담당업무
- 우대사항
- 전형절차
- 필요한 역량
- 기타 요구사항
"""
)


# --------------------------------------------------
# 3️⃣ 모드별 분기: 기존 자소서 입력 vs 생성할 항목 선택
# --------------------------------------------------

resume_text = ""
selected_items = []

ITEM_PRESETS = [
    "지원동기",
    "성장과정",
    "성격의 장점과 단점",
    "직무 관련 역량 및 전문성",
    "입사 후 포부",
    "협업 및 팀워크 경험",
    "위기 극복 / 문제해결 경험",
    "윤리의식과 책임감",
    "자기계발 노력",
    "봉사 및 헌신 경험",
]

if mode == "🔍 기존 자소서 분석 및 수정":

    st.header("3️⃣ 자기소개서")

    input_mode = st.radio(
        "입력 방식을 선택하세요",
        ["통째로 입력 / 파일 업로드", "문항별로 나눠서 입력"],
        horizontal=True
    )

    if input_mode == "통째로 입력 / 파일 업로드":

        uploaded_file = st.file_uploader(
            "📎 자기소개서 파일 업로드",
            type=["pdf", "docx", "txt"]
        )

        if uploaded_file is not None:

            resume_text = extract_text(uploaded_file)

            if resume_text:
                st.success(f"파일을 읽었습니다: {uploaded_file.name}")
                with st.expander("📖 추출된 자기소개서 확인"):
                    st.text_area("파일 내용", resume_text, height=300, key="extracted_preview")
            else:
                st.error("파일에서 텍스트를 읽지 못했습니다.")

        direct_text = st.text_area(
            "또는 여기에 자기소개서를 직접 입력하세요",
            height=200
        )

        if not resume_text and direct_text.strip():
            resume_text = direct_text
            st.info("✍️ 직접 입력한 자기소개서를 분석에 사용합니다.")

    else:
        st.caption("문항 제목과 답변을 나눠서 입력하면, AI가 각 문항의 맥락을 더 명확히 구분해서 분석합니다.")

        if "qa_sections" not in st.session_state:
            st.session_state.qa_sections = [
                {"title": "", "body": ""}
            ]

        to_remove = None

        for idx, section in enumerate(st.session_state.qa_sections):
            with st.container(border=True):
                c1, c2 = st.columns([5, 1])
                with c1:
                    section["title"] = st.text_input(
                        f"문항 {idx + 1} 제목",
                        value=section["title"],
                        placeholder="예: 지원 동기를 작성해 주세요",
                        key=f"qa_title_{idx}"
                    )
                with c2:
                    st.write("")
                    st.write("")
                    if len(st.session_state.qa_sections) > 1:
                        if st.button("삭제", key=f"qa_remove_{idx}"):
                            to_remove = idx

                section["body"] = st.text_area(
                    f"문항 {idx + 1} 답변",
                    value=section["body"],
                    height=150,
                    key=f"qa_body_{idx}"
                )
                st.caption(f"글자 수: {count_chars(section['body'])}자")

        if to_remove is not None:
            st.session_state.qa_sections.pop(to_remove)
            st.rerun()

        if st.button("➕ 문항 추가"):
            st.session_state.qa_sections.append({"title": "", "body": ""})
            st.rerun()

        combined_parts = []
        for idx, section in enumerate(st.session_state.qa_sections):
            title = section["title"].strip() or f"문항 {idx + 1}"
            body = section["body"].strip()
            if body:
                combined_parts.append(f"[{title}]\n{body}")

        resume_text = "\n\n".join(combined_parts)

else:
    # ✨ 새 자소서 항목별로 생성 모드
    st.header("3️⃣ 작성할 자기소개서 항목 선택")
    st.caption("지원하는 병원(기업)에서 요구하는 항목을 선택하세요. 목록에 없으면 직접 추가할 수 있어요.")

    selected_presets = st.multiselect(
        "📋 자소서 항목 선택",
        ITEM_PRESETS,
        default=["지원동기", "성장과정", "입사 후 포부"]
    )

    custom_items_raw = st.text_input(
        "➕ 목록에 없는 항목 직접 추가 (쉼표로 구분)",
        placeholder="예: 존경하는 인물과 그 이유, 간호관, 인생에서 가장 힘들었던 순간"
    )
    custom_items = [x.strip() for x in custom_items_raw.split(",") if x.strip()]

    selected_items = selected_presets + custom_items

    if selected_items:
        st.success(f"선택된 항목 ({len(selected_items)}개): " + ", ".join(selected_items))
    else:
        st.warning("최소 1개 이상의 항목을 선택하거나 추가해주세요.")


st.header("4️⃣ 직무 관련 최신 이슈 (선택사항)")

st.caption("직무나 업계 관련 최신 뉴스를 가져와서, 분석/생성 시 배경 참고 자료로 함께 활용할 수 있습니다.")

news_col1, news_col2 = st.columns([3, 1])

with news_col1:
    news_keyword = st.text_input(
        "🔍 검색 키워드",
        value=job if job else "",
        placeholder="예: 간호사, 신규간호사 처우, 병원 채용 등"
    )

with news_col2:
    st.write("")
    st.write("")
    fetch_news_button = st.button("📰 뉴스 불러오기", use_container_width=True)

if fetch_news_button:
    if not news_keyword.strip():
        st.warning("검색 키워드를 입력해주세요.")
    else:
        try:
            with st.spinner("최신 뉴스를 가져오는 중..."):
                news_items = fetch_related_news(news_keyword)
            st.session_state["news_items"] = news_items
            st.session_state["news_keyword"] = news_keyword
            if not news_items:
                st.info("관련 뉴스를 찾지 못했습니다. 다른 키워드로 시도해보세요.")
        except Exception as e:
            st.error("뉴스를 가져오는 중 오류가 발생했습니다.")
            st.code(str(e))

if "news_items" in st.session_state and st.session_state["news_items"]:
    with st.expander(f"📰 '{st.session_state.get('news_keyword','')}' 관련 최신 뉴스 ({len(st.session_state['news_items'])}건)", expanded=True):
        for n in st.session_state["news_items"]:
            st.markdown(f"- [{n['title']}]({n['link']})  \n  <sub>{n['source']} · {n['pub_date']}</sub>", unsafe_allow_html=True)

    use_news_in_analysis = st.checkbox(
        "✅ 분석/생성할 때 위 뉴스 헤드라인을 배경 참고 자료로 함께 활용하기",
        value=True
    )
else:
    use_news_in_analysis = False


st.header("5️⃣ 나의 추가 정보")

profile = st.text_area(
    "👤 경력 / 학력 / 자격증 / 경험 / 스펙",
    height=180,
    placeholder="""예:

- 간호학과 졸업
- 대학병원 실습 경험
- BLS 자격증
- 봉사활동
- 아르바이트
- 팀 프로젝트 경험
- 외국어
- 기타 본인이 가지고 있는 경험
"""
)


# --------------------------------------------------
# 글자 수 표시 + 목표 글자 수
# --------------------------------------------------

st.divider()

if mode == "🔍 기존 자소서 분석 및 수정" and resume_text:
    st.caption(f"📏 현재 자기소개서 글자 수 (공백 포함): **{count_chars(resume_text)}자**")

length_label = (
    "📏 목표 글자 수 (공백 포함, 전체 기준)"
    if mode == "🔍 기존 자소서 분석 및 수정"
    else "📏 항목당 목표 글자 수 (공백 포함)"
)

max_length_option = st.selectbox(
    length_label,
    ["제한 없음", "300자 내외", "500자 내외", "700자 내외", "1000자 내외", "직접 입력"]
)

target_length = ""
if max_length_option == "직접 입력":
    custom_len = st.number_input("목표 글자 수를 입력하세요 (자)", min_value=100, max_value=3000, value=500, step=50)
    target_length = f"약 {custom_len}자"
elif max_length_option != "제한 없음":
    target_length = max_length_option

if mode == "🔍 기존 자소서 분석 및 수정" and target_length and resume_text:
    diff = None
    if max_length_option == "직접 입력":
        diff = count_chars(resume_text) - custom_len
    if diff is not None:
        if diff > 0:
            st.warning(f"목표보다 {diff}자 초과되었습니다.")
        elif diff < 0:
            st.info(f"목표보다 {abs(diff)}자 여유가 있습니다.")


# ====================================================
# 모드 1: 기존 자소서 분석 및 수정 (기존 로직)
# ====================================================

if mode == "🔍 기존 자소서 분석 및 수정":

    analyze_button = st.button(
        "🔍 종합 분석 시작",
        type="primary",
        use_container_width=True
    )

    if analyze_button:

        if not api_key:
            st.error("먼저 왼쪽 사이드바에 API Key를 입력해주세요.")
            st.stop()

        if not company:
            st.warning("지원 회사를 입력해주세요.")
            st.stop()

        if not job:
            st.warning("지원 직무를 입력해주세요.")
            st.stop()

        if not resume_text:
            st.warning("자기소개서를 입력하거나 업로드해주세요.")
            st.stop()

        if not job_posting:
            st.warning("채용공고 내용을 입력해주세요.")
            st.stop()

        try:
            with st.spinner(f"{provider}가 자기소개서와 채용공고를 분석하고 있습니다..."):
                analysis = analyze_resume(
                    provider,
                    api_key,
                    model_name,
                    resume_text,
                    company,
                    job,
                    job_posting,
                    profile,
                    company_philosophy,
                    reference_resume,
                    target_length,
                    st.session_state.get("news_items") if use_news_in_analysis else None
                )

            st.session_state["analysis"] = analysis
            st.session_state["resume_text"] = resume_text
            st.session_state["company"] = company
            st.session_state["job"] = job
            st.session_state["target_length"] = target_length
            st.session_state["provider"] = provider
            st.session_state["model_name"] = model_name

            st.success("분석이 완료되었습니다.")

        except Exception as e:
            st.error("AI 분석 중 오류가 발생했습니다.")
            st.code(str(e))

    # 분석 결과
    if "analysis" in st.session_state:

        st.divider()
        st.header("📊 AI 종합 분석")
        st.caption(f"사용된 AI: {st.session_state.get('provider', '')} ({st.session_state.get('model_name', '')})")
        st.markdown(st.session_state["analysis"])

        dl_col1, dl_col2 = st.columns(2)
        with dl_col1:
            st.download_button(
                "⬇️ 분석 결과 다운로드 (.txt)",
                data=st.session_state["analysis"],
                file_name=f"{st.session_state.get('company','분석결과')}_자소서분석.txt",
                mime="text/plain",
                use_container_width=True
            )
        with dl_col2:
            docx_bytes = build_docx_bytes("자기소개서 분석 결과", st.session_state["analysis"])
            if docx_bytes:
                st.download_button(
                    "⬇️ 분석 결과 다운로드 (.docx)",
                    data=docx_bytes,
                    file_name=f"{st.session_state.get('company','분석결과')}_자소서분석.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    use_container_width=True
                )

        # 수정본 생성
        st.divider()
        st.header("6️⃣ AI 자기소개서 수정")

        st.write(
            "분석 결과를 바탕으로 기존 경험을 유지하면서 "
            "지원 직무에 맞게 자기소개서를 개선합니다."
        )

        revise_button = st.button(
            "✏️ AI 수정본 만들기",
            type="primary",
            use_container_width=True
        )

        if revise_button:

            if not api_key:
                st.error("왼쪽 사이드바에 API Key를 입력해주세요.")
                st.stop()

            try:
                with st.spinner(f"{provider}가 자기소개서를 수정하고 있습니다..."):
                    revised = revise_resume(
                        provider,
                        api_key,
                        model_name,
                        st.session_state["resume_text"],
                        st.session_state["analysis"],
                        st.session_state["company"],
                        st.session_state["job"],
                        st.session_state.get("target_length", "")
                    )

                st.session_state["revised"] = revised

            except Exception as e:
                st.error("수정본 생성 중 오류가 발생했습니다.")
                st.code(str(e))

    # 원본 / 수정본 비교
    if "revised" in st.session_state:

        st.divider()
        st.header("🔄 원본 vs AI 수정본")

        left, right = st.columns(2)

        with left:
            st.subheader("📄 원본")
            st.caption(f"글자 수: {count_chars(st.session_state['resume_text'])}자")
            st.text_area(
                "원본 자기소개서",
                st.session_state["resume_text"],
                height=600,
                key="original_text"
            )

        with right:
            st.subheader("✨ AI 수정본")
            st.caption(f"글자 수(수정사항 설명 포함): {count_chars(st.session_state['revised'])}자")
            st.text_area(
                "AI 수정본 및 수정 설명",
                st.session_state["revised"],
                height=600,
                key="revised_text"
            )

        dl_col1, dl_col2 = st.columns(2)
        with dl_col1:
            st.download_button(
                "⬇️ 수정본 다운로드 (.txt)",
                data=st.session_state["revised"],
                file_name=f"{st.session_state.get('company','수정본')}_자소서수정본.txt",
                mime="text/plain",
                use_container_width=True
            )
        with dl_col2:
            revised_docx = build_docx_bytes("AI 자기소개서 수정본", st.session_state["revised"])
            if revised_docx:
                st.download_button(
                    "⬇️ 수정본 다운로드 (.docx)",
                    data=revised_docx,
                    file_name=f"{st.session_state.get('company','수정본')}_자소서수정본.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    use_container_width=True
                )

        st.info("💡 현재 버전에서는 AI 수정본 안에 '주요 수정사항'도 함께 표시됩니다.")


# ====================================================
# 모드 2: 새 자소서 항목별로 생성 (신규)
# ====================================================

else:

    generate_button = st.button(
        "✨ 자소서 항목별로 생성하기",
        type="primary",
        use_container_width=True
    )

    if generate_button:

        if not api_key:
            st.error("먼저 왼쪽 사이드바에 API Key를 입력해주세요.")
            st.stop()

        if not company:
            st.warning("지원 회사를 입력해주세요.")
            st.stop()

        if not job:
            st.warning("지원 직무를 입력해주세요.")
            st.stop()

        if not profile or not profile.strip():
            st.warning("자소서에 담을 실제 경험/스펙 정보를 '5️⃣ 나의 추가 정보'에 입력해주세요. (없으면 AI가 내용을 지어낼 수밖에 없어요)")
            st.stop()

        if not selected_items:
            st.warning("작성할 자소서 항목을 최소 1개 이상 선택하거나 추가해주세요.")
            st.stop()

        try:
            with st.spinner(f"{provider}가 선택하신 {len(selected_items)}개 항목의 자소서를 작성하고 있습니다..."):
                generated = generate_resume(
                    provider,
                    api_key,
                    model_name,
                    company,
                    job,
                    job_posting,
                    profile,
                    company_philosophy,
                    reference_resume,
                    selected_items,
                    target_length,
                    st.session_state.get("news_items") if use_news_in_analysis else None
                )

            st.session_state["generated"] = generated
            st.session_state["gen_company"] = company
            st.session_state["gen_job"] = job
            st.session_state["gen_provider"] = provider
            st.session_state["gen_model_name"] = model_name

            st.success("자소서 생성이 완료되었습니다.")

        except Exception as e:
            st.error("AI 생성 중 오류가 발생했습니다.")
            st.code(str(e))

    if "generated" in st.session_state:

        st.divider()
        st.header("✨ AI가 생성한 자기소개서")
        st.caption(f"사용된 AI: {st.session_state.get('gen_provider', '')} ({st.session_state.get('gen_model_name', '')})")
        st.markdown(st.session_state["generated"])

        dl_col1, dl_col2 = st.columns(2)
        with dl_col1:
            st.download_button(
                "⬇️ 생성 결과 다운로드 (.txt)",
                data=st.session_state["generated"],
                file_name=f"{st.session_state.get('gen_company','자소서')}_AI생성자소서.txt",
                mime="text/plain",
                use_container_width=True
            )
        with dl_col2:
            gen_docx_bytes = build_docx_bytes("AI 생성 자기소개서", st.session_state["generated"])
            if gen_docx_bytes:
                st.download_button(
                    "⬇️ 생성 결과 다운로드 (.docx)",
                    data=gen_docx_bytes,
                    file_name=f"{st.session_state.get('gen_company','자소서')}_AI생성자소서.docx",
                    mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                    use_container_width=True
                )

        st.info("💡 (※ 보완 필요) 표시가 있는 부분은, 반드시 본인의 실제 구체적 경험/수치로 채워 넣은 뒤 제출하세요. AI가 지어낸 내용이 아니라 사실 확인이 꼭 필요합니다.")


# --------------------------------------------------
# 하단 안내
# --------------------------------------------------

st.divider()

st.caption(
    "※ AI 평가/생성 결과는 실제 채용 결과나 합격 확률을 보장하지 않습니다. "
    "특히 존재하지 않는 경력·성과·자격증을 AI가 만들어내지 않도록 설계되어 있지만, "
    "제출 전 반드시 본인이 직접 검토하고 사실 확인을 해주세요."
)
